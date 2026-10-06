#!/usr/bin/env python3
"""
RVA amino-acid substitution analysis: Grantham-distance support for a
binary (physicochemical class) classification of amino-acid substitutions.

HOW TO USE
----------
1. Put one input file per gene/protein (.xlsx or .csv) in a folder, e.g. data/VP1.xlsx, data/VP2.xlsx ...
   Each file needs (column names are detected automatically, see COLUMN_MAP to override):
        Position | Amino acid change  (e.g. "M -> K")  | Count
   or   Position | Ref | Alt | Count
   Repeated codon rows are fine: they are collapsed to unique Position+Ref+Alt.
2. Run:  python rva_substitution_analysis.py
   An upload / file-selection box opens: choose your gene file(s).
   (Alternatively paste a folder/file path into INPUT_PATH, or use --input PATH --outdir PATH.
    In Google Colab the box is the Colab upload button and the results are downloaded as a zip.)
3. Results appear in results/<GENE>/ and results/ALL_GENES_summary.xlsx next to your data.

Requires: pandas, numpy, scipy, matplotlib, openpyxl   (pip install pandas numpy scipy matplotlib openpyxl)
"""

# ╔══════════════════════════════════════════════════════════════════════════╗
# ║                         >>>  YOUR INPUT  <<<                             ║
# ║  Leave INPUT_PATH empty ("")  ->  an UPLOAD / file-selection box opens    ║
# ║  when you run the script: select one or many files (one file per gene,    ║
# ║  .xlsx or .csv; hold Ctrl/Shift to select several).                       ║
# ║  Or paste a folder / file path instead, e.g. "C:/Users/me/Desktop/data"   ║
# ║  Results are saved in a new folder called "results" next to your data.    ║
# ╚══════════════════════════════════════════════════════════════════════════╝
INPUT_PATH = ""
# ════════════════════════════════════════════════════════════════════════════

import argparse, os, re, sys, math
import numpy as np
import pandas as pd
from scipy import stats
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.colors import ListedColormap, LogNorm
from matplotlib.patches import Patch

# =====================  SETTINGS YOU MAY EDIT  =====================
# Physicochemical classes (conservative = same class, radical = different class)
CLASSES = {
    "Hydrophobic": "AVLIMFYWGP",
    "Polar":       "STNQC",
    "Basic":       "HRK",
    "Acidic":      "DE",
}
# Grantham categories (Li et al. 1984): upper limits of each band
GRANTHAM_CUTOFFS = [(50, "Conservative"), (100, "Moderately conservative"),
                    (150, "Moderately radical"), (float("inf"), "Radical")]
BINARY_CUTOFF_GRANTHAM = 50      # <= 50 -> "conservative" when turning Grantham into a binary call
# If automatic column detection fails, set names here, e.g. {"position": "Pos", "change": "AA change", "count": "N"}
COLUMN_MAP = {}   # keys: position, change, ref, alt, count
# ===================================================================

# ---- published Grantham matrix (embedded so the script needs no downloads) ----
# Grantham R (1974) Science 185:862-864 (values taken from AAindex entry GRAR740104 - please verify against the paper.)
GRANTHAM = {
    'A': {'A': 0, 'R': 112, 'N': 111, 'D': 126, 'C': 195, 'Q': 91, 'E': 107, 'G': 60, 'H': 86, 'I': 94, 'L': 96, 'K': 106, 'M': 84, 'F': 113, 'P': 27, 'S': 99, 'T': 58, 'W': 148, 'Y': 112, 'V': 64},
    'R': {'A': 112, 'R': 0, 'N': 86, 'D': 96, 'C': 180, 'Q': 43, 'E': 54, 'G': 125, 'H': 29, 'I': 97, 'L': 102, 'K': 26, 'M': 91, 'F': 97, 'P': 103, 'S': 110, 'T': 71, 'W': 101, 'Y': 77, 'V': 96},
    'N': {'A': 111, 'R': 86, 'N': 0, 'D': 23, 'C': 139, 'Q': 46, 'E': 42, 'G': 80, 'H': 68, 'I': 149, 'L': 153, 'K': 94, 'M': 142, 'F': 158, 'P': 91, 'S': 46, 'T': 65, 'W': 174, 'Y': 143, 'V': 133},
    'D': {'A': 126, 'R': 96, 'N': 23, 'D': 0, 'C': 154, 'Q': 61, 'E': 45, 'G': 94, 'H': 81, 'I': 168, 'L': 172, 'K': 101, 'M': 160, 'F': 177, 'P': 108, 'S': 65, 'T': 85, 'W': 181, 'Y': 160, 'V': 152},
    'C': {'A': 195, 'R': 180, 'N': 139, 'D': 154, 'C': 0, 'Q': 154, 'E': 170, 'G': 159, 'H': 174, 'I': 198, 'L': 198, 'K': 202, 'M': 196, 'F': 205, 'P': 169, 'S': 112, 'T': 149, 'W': 215, 'Y': 194, 'V': 192},
    'Q': {'A': 91, 'R': 43, 'N': 46, 'D': 61, 'C': 154, 'Q': 0, 'E': 29, 'G': 87, 'H': 24, 'I': 109, 'L': 113, 'K': 53, 'M': 101, 'F': 116, 'P': 76, 'S': 68, 'T': 42, 'W': 130, 'Y': 99, 'V': 96},
    'E': {'A': 107, 'R': 54, 'N': 42, 'D': 45, 'C': 170, 'Q': 29, 'E': 0, 'G': 98, 'H': 40, 'I': 134, 'L': 138, 'K': 56, 'M': 126, 'F': 140, 'P': 93, 'S': 80, 'T': 65, 'W': 152, 'Y': 122, 'V': 121},
    'G': {'A': 60, 'R': 125, 'N': 80, 'D': 94, 'C': 159, 'Q': 87, 'E': 98, 'G': 0, 'H': 98, 'I': 135, 'L': 138, 'K': 127, 'M': 127, 'F': 153, 'P': 42, 'S': 56, 'T': 59, 'W': 184, 'Y': 147, 'V': 109},
    'H': {'A': 86, 'R': 29, 'N': 68, 'D': 81, 'C': 174, 'Q': 24, 'E': 40, 'G': 98, 'H': 0, 'I': 94, 'L': 99, 'K': 32, 'M': 87, 'F': 100, 'P': 77, 'S': 89, 'T': 47, 'W': 115, 'Y': 83, 'V': 84},
    'I': {'A': 94, 'R': 97, 'N': 149, 'D': 168, 'C': 198, 'Q': 109, 'E': 134, 'G': 135, 'H': 94, 'I': 0, 'L': 5, 'K': 102, 'M': 10, 'F': 21, 'P': 95, 'S': 142, 'T': 89, 'W': 61, 'Y': 33, 'V': 29},
    'L': {'A': 96, 'R': 102, 'N': 153, 'D': 172, 'C': 198, 'Q': 113, 'E': 138, 'G': 138, 'H': 99, 'I': 5, 'L': 0, 'K': 107, 'M': 15, 'F': 22, 'P': 98, 'S': 145, 'T': 92, 'W': 61, 'Y': 36, 'V': 32},
    'K': {'A': 106, 'R': 26, 'N': 94, 'D': 101, 'C': 202, 'Q': 53, 'E': 56, 'G': 127, 'H': 32, 'I': 102, 'L': 107, 'K': 0, 'M': 95, 'F': 102, 'P': 103, 'S': 121, 'T': 78, 'W': 110, 'Y': 85, 'V': 97},
    'M': {'A': 84, 'R': 91, 'N': 142, 'D': 160, 'C': 196, 'Q': 101, 'E': 126, 'G': 127, 'H': 87, 'I': 10, 'L': 15, 'K': 95, 'M': 0, 'F': 28, 'P': 87, 'S': 135, 'T': 81, 'W': 67, 'Y': 36, 'V': 21},
    'F': {'A': 113, 'R': 97, 'N': 158, 'D': 177, 'C': 205, 'Q': 116, 'E': 140, 'G': 153, 'H': 100, 'I': 21, 'L': 22, 'K': 102, 'M': 28, 'F': 0, 'P': 114, 'S': 155, 'T': 103, 'W': 40, 'Y': 22, 'V': 50},
    'P': {'A': 27, 'R': 103, 'N': 91, 'D': 108, 'C': 169, 'Q': 76, 'E': 93, 'G': 42, 'H': 77, 'I': 95, 'L': 98, 'K': 103, 'M': 87, 'F': 114, 'P': 0, 'S': 74, 'T': 38, 'W': 147, 'Y': 110, 'V': 68},
    'S': {'A': 99, 'R': 110, 'N': 46, 'D': 65, 'C': 112, 'Q': 68, 'E': 80, 'G': 56, 'H': 89, 'I': 142, 'L': 145, 'K': 121, 'M': 135, 'F': 155, 'P': 74, 'S': 0, 'T': 58, 'W': 177, 'Y': 144, 'V': 124},
    'T': {'A': 58, 'R': 71, 'N': 65, 'D': 85, 'C': 149, 'Q': 42, 'E': 65, 'G': 59, 'H': 47, 'I': 89, 'L': 92, 'K': 78, 'M': 81, 'F': 103, 'P': 38, 'S': 58, 'T': 0, 'W': 128, 'Y': 92, 'V': 69},
    'W': {'A': 148, 'R': 101, 'N': 174, 'D': 181, 'C': 215, 'Q': 130, 'E': 152, 'G': 184, 'H': 115, 'I': 61, 'L': 61, 'K': 110, 'M': 67, 'F': 40, 'P': 147, 'S': 177, 'T': 128, 'W': 0, 'Y': 37, 'V': 88},
    'Y': {'A': 112, 'R': 77, 'N': 143, 'D': 160, 'C': 194, 'Q': 99, 'E': 122, 'G': 147, 'H': 83, 'I': 33, 'L': 36, 'K': 85, 'M': 36, 'F': 22, 'P': 110, 'S': 144, 'T': 92, 'W': 37, 'Y': 0, 'V': 55},
    'V': {'A': 64, 'R': 96, 'N': 133, 'D': 152, 'C': 192, 'Q': 96, 'E': 121, 'G': 109, 'H': 84, 'I': 29, 'L': 32, 'K': 97, 'M': 21, 'F': 50, 'P': 68, 'S': 124, 'T': 69, 'W': 88, 'Y': 55, 'V': 0},
}


THREE = {"ALA":"A","ARG":"R","ASN":"N","ASP":"D","CYS":"C","GLN":"Q","GLU":"E","GLY":"G","HIS":"H","ILE":"I",
         "LEU":"L","LYS":"K","MET":"M","PHE":"F","PRO":"P","SER":"S","THR":"T","TRP":"W","TYR":"Y","VAL":"V"}
AA20 = "ARNDCQEGHILKMFPSTWYV"
AA_CLASS = {aa: cls for cls, aas in CLASSES.items() for aa in aas}
assert set(AA_CLASS) == set(AA20), "CLASSES must cover each of the 20 amino acids exactly once"
ORDER = [aa for cls in CLASSES for aa in CLASSES[cls]]            # heatmap order, grouped by class
CATS = [c for _, c in GRANTHAM_CUTOFFS]
CAT_COLORS = ["#2e7d32", "#9ccc65", "#ffb74d", "#c62828"]


def gcat(d):
    for lim, name in GRANTHAM_CUTOFFS:
        if d <= lim:
            return name


# ------------------------------ input reading ------------------------------
def find_col(df, key, candidates, contains=False):
    if COLUMN_MAP.get(key):
        return COLUMN_MAP[key]
    low = {str(c).strip().lower(): c for c in df.columns}
    for cand in candidates:
        if cand in low:
            return low[cand]
    if contains:
        for name, c in low.items():
            if any(cand in name for cand in candidates):
                return c
    return None


def parse_aa(token):
    t = str(token).strip()
    if len(t) == 1:
        return t.upper()
    return THREE.get(t.upper(), t.upper())


def read_gene_table(df, gene):
    pos = find_col(df, "position", ["position", "pos", "site", "residue"])
    cnt = find_col(df, "count", ["count", "frequency", "freq", "n"])
    chg = find_col(df, "change", ["amino acid change", "aa change", "substitution", "aa_change", "change"], contains=True)
    ref = find_col(df, "ref", ["ref", "ref_aa", "reference", "consensus", "from"])
    alt = find_col(df, "alt", ["alt", "alt_aa", "substituted", "alternative", "mutant", "to"])
    if pos is None:
        sys.exit(f"[{gene}] no Position column found. Columns: {list(df.columns)}. Set COLUMN_MAP.")
    if ref is not None and alt is not None:
        r = df[ref].map(parse_aa); a = df[alt].map(parse_aa)
    elif chg is not None:
        pat = re.compile(r"([A-Za-z\*]+)\s*(?:->|→|=>|>)\s*([A-Za-z\*]+)")
        m = df[chg].astype(str).str.extract(pat)
        r = m[0].map(lambda x: parse_aa(x) if isinstance(x, str) else None)
        a = m[1].map(lambda x: parse_aa(x) if isinstance(x, str) else None)
    else:
        sys.exit(f"[{gene}] need either an 'Amino acid change' column or Ref+Alt columns. Columns: {list(df.columns)}")
    out = pd.DataFrame({"Position": df[pos], "Ref": r, "Alt": a})
    out["Count"] = pd.to_numeric(df[cnt], errors="coerce").fillna(1) if cnt is not None else 1
    out["_row"] = range(len(df))
    return out, df


def clean_and_collapse(t, gene):
    n_rows = len(t)
    bad = ~(t.Ref.isin(list(AA20)) & t.Alt.isin(list(AA20)))
    n_bad = int(bad.sum())
    t = t[~bad]
    syn = t.Ref == t.Alt
    n_syn = int(syn.sum())
    t = t[~syn]
    grp = t.groupby(["Position", "Ref", "Alt"], sort=False)
    u = grp.agg(Count=("Count", "first"), n_rows=("Count", "size"), cmin=("Count", "min"), cmax=("Count", "max")).reset_index()
    varying = int((u.cmin != u.cmax).sum())
    if varying:
        print(f"  [{gene}] WARNING: Count differs between rows of the same substitution in {varying} cases; the first value is used.")
    u = u.drop(columns=["cmin", "cmax"])
    info = dict(rows_in=n_rows, rows_dropped_invalid=n_bad, rows_dropped_synonymous=n_syn, unique_substitutions=len(u),
                positions=u.Position.nunique())
    return u, info


# ------------------------------ scoring ------------------------------
def add_scores(u):
    u = u.copy()
    u["Grantham"] = [GRANTHAM[a][b] for a, b in zip(u.Ref, u.Alt)]
    u["Grantham_category"] = u.Grantham.map(gcat)
    u["Ref_class"] = u.Ref.map(AA_CLASS); u["Alt_class"] = u.Alt.map(AA_CLASS)
    u["Class_change"] = np.where(u.Ref_class == u.Alt_class, "Conservative", "Radical")
    return u


def kappa(a, b):
    a = np.asarray(a); b = np.asarray(b)
    po = (a == b).mean()
    pe = sum(((a == k).mean()) * ((b == k).mean()) for k in (0, 1))
    return (po - pe) / (1 - pe) if pe < 1 else float("nan")


def auc_from_u(cons, rad):
    """P(radical value > conservative value); ties count 0.5"""
    U = stats.mannwhitneyu(rad, cons, alternative="two-sided")
    return U.statistic / (len(cons) * len(rad)), U.pvalue


def compare_metric(u, col, sign=1):
    c = u.loc[u.Class_change == "Conservative", col]; r = u.loc[u.Class_change == "Radical", col]
    if len(c) < 2 or len(r) < 2:
        return dict(metric=col)
    auc, p = auc_from_u(sign * c, sign * r)
    w = u.Count
    return dict(metric=col, cons_mean=c.mean(), cons_median=c.median(), rad_mean=r.mean(), rad_median=r.median(),
                cons_wmean=np.average(c, weights=w[c.index]), rad_wmean=np.average(r, weights=w[r.index]),
                mannwhitney_p=p, AUC=auc)


def agreement(u, pred_radical):
    y = (u.Class_change == "Radical").astype(int).values
    p = np.asarray(pred_radical).astype(int)
    return dict(agreement_pct=100 * (p == y).mean(),
                agreement_weighted_pct=100 * np.average(p == y, weights=u.Count),
                kappa=kappa(y, p))


def analyse(u):
    res = {}
    res["metrics"] = pd.DataFrame([compare_metric(u, "Grantham", 1)])
    ag = {"Grantham (<=%g conservative)" % BINARY_CUTOFF_GRANTHAM: agreement(u, u.Grantham > BINARY_CUTOFF_GRANTHAM)}
    res["agreement"] = pd.DataFrame(ag).T.reset_index().rename(columns={"index": "criterion"})
    # sensitivity to Grantham cut-off
    rows = []
    for t in (30, 40, 50, 60, 70, 75, 80, 100):
        d = agreement(u, u.Grantham > t); d["Grantham_cutoff"] = t; rows.append(d)
    res["cutoff_sensitivity"] = pd.DataFrame(rows)[["Grantham_cutoff", "agreement_pct", "agreement_weighted_pct", "kappa"]]
    # four-category cross-tab
    ct = pd.crosstab(u.Grantham_category, u.Class_change).reindex(index=CATS, columns=["Conservative", "Radical"], fill_value=0)
    ct.columns = ["Class-conservative (n)", "Class-radical (n)"]
    ct.insert(0, "n", ct.sum(axis=1))
    ct["Class-conservative (%)"] = 100 * ct.iloc[:, 1] / ct.n.replace(0, np.nan)
    ct["Class-radical (%)"] = 100 * ct.iloc[:, 2] / ct.n.replace(0, np.nan)
    ct["% of all unique"] = 100 * ct.n / ct.n.sum()
    wc = u.groupby("Grantham_category").Count.sum().reindex(CATS).fillna(0)
    ct["Total count"] = wc.values; ct["% of total count"] = 100 * wc.values / wc.sum()
    res["crosstab"] = ct.reset_index().rename(columns={"Grantham_category": "Grantham category"})
    # discordant substitution types
    u2 = u.assign(Substitution=u.Ref + " > " + u.Alt)
    sub = u2.groupby("Substitution").agg(Ref=("Ref", "first"), Alt=("Alt", "first"), n_positions=("Position", "nunique"),
                                         total_count=("Count", "sum"), Grantham=("Grantham", "first"),
                                         Grantham_category=("Grantham_category", "first"), Ref_class=("Ref_class", "first"),
                                         Alt_class=("Alt_class", "first"), Class_change=("Class_change", "first")).reset_index()
    sub = sub.sort_values(["total_count", "n_positions", "Grantham"], ascending=[False, False, True]).reset_index(drop=True)
    sub.insert(0, "Rank", sub.index + 1)
    res["ranked"] = sub
    res["discord_cons_but_G_gt"] = sub[(sub.Class_change == "Conservative") & (sub.Grantham > BINARY_CUTOFF_GRANTHAM)]
    res["discord_rad_but_G_le"] = sub[(sub.Class_change == "Radical") & (sub.Grantham <= BINARY_CUTOFF_GRANTHAM)]
    # per-position summary
    pp = u.groupby("Position").agg(Ref=("Ref", "first"), n_alt_aa=("Alt", "nunique"), total_count=("Count", "sum"),
                                   mean_Grantham=("Grantham", "mean"), min_Grantham=("Grantham", "min"),
                                   max_Grantham=("Grantham", "max")).reset_index()
    pp["Grantham_range"] = pp.max_Grantham - pp.min_Grantham
    res["positions"] = pp
    # frequency vs severity
    if len(u) > 3:
        rho, p = stats.spearmanr(u.Grantham, u.Count)
        res["spearman_G_vs_count"] = (rho, p)
    return res


# ------------------------------ matrices / figures ------------------------------
def matrices(u):
    cnt = pd.DataFrame(0, index=ORDER, columns=ORDER, dtype=float)
    pos = pd.DataFrame(0, index=ORDER, columns=ORDER, dtype=float)
    for r in u.itertuples():
        cnt.loc[r.Ref, r.Alt] += r.Count; pos.loc[r.Ref, r.Alt] += 1
    gr = pd.DataFrame([[GRANTHAM[a][b] for b in ORDER] for a in ORDER], index=ORDER, columns=ORDER)
    return cnt, pos, gr


def class_blocks(ax):
    b = 0
    for cls in list(CLASSES)[:-1]:
        b += len(CLASSES[cls]); ax.axhline(b - .5, c="k", lw=1); ax.axvline(b - .5, c="k", lw=1)


def draw_figures(u, cnt, gene, outdir):
    n = len(ORDER); idx = {a: i for i, a in enumerate(ORDER)}
    catm = np.full((n, n), np.nan); gm = np.full((n, n), np.nan)
    for r in u.drop_duplicates(["Ref", "Alt"]).itertuples():
        i, j = idx[r.Ref], idx[r.Alt]; gm[i, j] = r.Grantham; catm[i, j] = CATS.index(r.Grantham_category)
    # Fig 1: Grantham category heatmap
    fig, ax = plt.subplots(figsize=(11, 9.5))
    ax.imshow(catm, cmap=ListedColormap(CAT_COLORS), vmin=-.5, vmax=3.5)
    for i in range(n):
        for j in range(n):
            if not np.isnan(gm[i, j]):
                ax.text(j, i, int(gm[i, j]), ha="center", va="center", fontsize=7,
                        color="white" if catm[i, j] in (0, 3) else "black")
    for a_, setter in ((ax.set_xticks, ax.set_xticklabels), (ax.set_yticks, ax.set_yticklabels)):
        a_(range(n)); setter(ORDER)
    ax.set_xlabel("Substituted amino acid"); ax.set_ylabel("Consensus (reference) amino acid")
    ax.xaxis.set_label_position("top"); ax.xaxis.tick_top(); class_blocks(ax)
    ax.set_title(f"{gene}: Grantham distance of observed substitutions (white = not observed)", pad=28)
    ax.legend(handles=[Patch(color=c, label=l) for c, l in zip(CAT_COLORS, ["Conservative (0-50)", "Mod. conservative (51-100)",
              "Mod. radical (101-150)", "Radical (>150)"])], loc="upper center", bbox_to_anchor=(.5, -.02), ncol=4, frameon=False, fontsize=8)
    plt.tight_layout(); plt.savefig(os.path.join(outdir, f"{gene}_Fig1_Grantham_heatmap.png"), dpi=200); plt.close()
    # Fig 2: frequency heatmap
    cm = cnt.values
    fig, ax = plt.subplots(figsize=(11, 9.5))
    im = ax.imshow(np.ma.masked_where(cm == 0, cm), cmap="YlOrRd", norm=LogNorm(vmin=1, vmax=max(cm.max(), 2)))
    for i in range(n):
        for j in range(n):
            if cm[i, j] > 0:
                ax.text(j, i, int(cm[i, j]), ha="center", va="center", fontsize=6.5, color="white" if cm[i, j] > cm.max() * .25 else "black")
    for a_, setter in ((ax.set_xticks, ax.set_xticklabels), (ax.set_yticks, ax.set_yticklabels)):
        a_(range(n)); setter(ORDER)
    ax.set_xlabel("Substituted amino acid"); ax.set_ylabel("Consensus (reference) amino acid")
    ax.xaxis.set_label_position("top"); ax.xaxis.tick_top(); class_blocks(ax)
    ax.set_title(f"{gene}: frequency of observed substitutions (log colour scale)", pad=28)
    plt.colorbar(im, fraction=.04, pad=.02, label="Total count (log)")
    plt.tight_layout(); plt.savefig(os.path.join(outdir, f"{gene}_Fig2_Count_heatmap.png"), dpi=200); plt.close()
    # Fig 3: category bars
    a = u.groupby("Grantham_category").size().reindex(CATS).fillna(0); b = u.groupby("Grantham_category").Count.sum().reindex(CATS).fillna(0)
    fig, ax = plt.subplots(1, 2, figsize=(11, 4.2))
    for k, (s, t) in enumerate([(a, f"Unique substitutions (n={int(a.sum())})"), (b, f"Weighted by Count (n={int(b.sum())})")]):
        bars = ax[k].bar(range(4), s.values, color=CAT_COLORS); ax[k].set_xticks(range(4))
        ax[k].set_xticklabels(["Conservative\n0-50", "Mod. cons.\n51-100", "Mod. radical\n101-150", "Radical\n>150"], fontsize=8)
        ax[k].set_title(t, fontsize=10); ax[k].set_ylabel("Number"); ax[k].set_ylim(0, max(s.max(), 1) * 1.15)
        for r_, v in zip(bars, s.values):
            ax[k].text(r_.get_x() + r_.get_width() / 2, v, f"{int(v)} ({100 * v / s.sum():.0f}%)", ha="center", va="bottom", fontsize=8)
    plt.suptitle(gene); plt.tight_layout(); plt.savefig(os.path.join(outdir, f"{gene}_Fig3_Grantham_categories.png"), dpi=200); plt.close()


# ------------------------------ text for the manuscript ------------------------------
def results_text(gene, info, res, u):
    ct = res["crosstab"]; g = res["metrics"].iloc[0]; ag = res["agreement"].iloc[0]
    def pct(i, col): return ct.loc[i, col]
    p = g.mannwhitney_p
    ptxt = "p < 0.001" if p < 0.001 else f"p = {p:.3f}"
    return (f"{gene}: {info['unique_substitutions']} unique substitutions at {info['positions']} positions.\n\n"
            f"Of the {int(ct.loc[0,'n'])} substitutions in the conservative Grantham category (0-50), {pct(0,'Class-conservative (%)'):.1f}% were also conservative by "
            f"physicochemical class and {pct(0,'Class-radical (%)'):.1f}% were radical by class. The proportion of class-radical substitutions was "
            f"{pct(1,'Class-radical (%)'):.1f}% in the moderately conservative category (51-100), {pct(2,'Class-radical (%)'):.1f}% in the moderately radical "
            f"category (101-150) and {pct(3,'Class-radical (%)'):.1f}% in the radical category (>150).\n"
            f"Overall, class-based and Grantham-based classification ({BINARY_CUTOFF_GRANTHAM} cut-off) agreed for {ag.agreement_pct:.1f}% of substitutions "
            f"({ag.agreement_weighted_pct:.1f}% when weighted by sequence count; Cohen's kappa = {ag.kappa:.2f}). Grantham distance was lower for conservative than radical "
            f"substitutions (mean {g.cons_mean:.1f} vs {g.rad_mean:.1f}; Mann-Whitney {ptxt}; AUC = {g.AUC:.2f}).\n")


# ------------------------------ Excel output ------------------------------
def style_sheet(ws, df_cols_widths=None):
    from openpyxl.styles import Font, PatternFill, Alignment
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F4E78")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "A2"
    for col in ws.columns:
        L = max(len(str(c.value)) if c.value is not None else 0 for c in list(col)[:200])
        ws.column_dimensions[col[0].column_letter].width = min(max(10, L + 2), 40)


def colour_category_column(ws, header="Grantham_category"):
    from openpyxl.styles import PatternFill
    fills = dict(zip(CATS, ["C8E6C9", "F0F4C3", "FFE0B2", "FFCDD2"]))
    heads = [c.value for c in ws[1]]
    if header not in heads: return
    j = heads.index(header) + 1
    for row in ws.iter_rows(min_row=2, min_col=j, max_col=j):
        for c in row:
            if c.value in fills: c.fill = PatternFill("solid", fgColor=fills[c.value])


def write_gene_workbook(path, gene, info, res, u, raw, cnt, pos, gr):
    from openpyxl.formatting.rule import ColorScaleRule
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        s = pd.DataFrame([info]).T.reset_index(); s.columns = ["Item", "Value"]; s.to_excel(xw, sheet_name="Summary", index=False)
        r0 = len(s) + 3
        res["crosstab"].round(2).to_excel(xw, sheet_name="Summary", startrow=r0, index=False)
        r1 = r0 + len(res["crosstab"]) + 3
        mm = res["metrics"].copy()
        mm[[c for c in mm.columns if c not in ("metric", "mannwhitney_p")]] = mm[[c for c in mm.columns if c not in ("metric", "mannwhitney_p")]].round(4)
        mm["mannwhitney_p"] = mm["mannwhitney_p"].map(lambda p: f"{p:.2e}")
        mm.to_excel(xw, sheet_name="Summary", startrow=r1, index=False)
        r2 = r1 + len(res["metrics"]) + 3
        res["agreement"].round(3).to_excel(xw, sheet_name="Summary", startrow=r2, index=False)
        r3 = r2 + len(res["agreement"]) + 3
        res["cutoff_sensitivity"].round(3).to_excel(xw, sheet_name="Summary", startrow=r3, index=False)
        ws = xw.sheets["Summary"]
        ws.cell(row=r0, column=1, value="Four Grantham categories vs binary (class) classification")
        ws.cell(row=r1, column=1, value="Class-conservative vs class-radical: score comparison (Mann-Whitney U; AUC = P(radical > conservative))")
        ws.cell(row=r2, column=1, value="Agreement of class-based classification with each distance measure")
        ws.cell(row=r3, column=1, value="Sensitivity of agreement to the Grantham cut-off")
        ws.column_dimensions["A"].width = 38
        res["ranked"].to_excel(xw, sheet_name="Substitutions_ranked", index=False)
        u.drop(columns=["n_rows"]).to_excel(xw, sheet_name="Unique_position_level", index=False)
        res["positions"].round(2).to_excel(xw, sheet_name="Position_summary", index=False)
        res["discord_cons_but_G_gt"].to_excel(xw, sheet_name="Class_cons_but_Grantham_high", index=False)
        res["discord_rad_but_G_le"].to_excel(xw, sheet_name="Class_rad_but_Grantham_low", index=False)
        raw.to_excel(xw, sheet_name="Data_with_Grantham", index=False)
        cnt.to_excel(xw, sheet_name="Matrix_count"); pos.to_excel(xw, sheet_name="Matrix_positions"); gr.to_excel(xw, sheet_name="Matrix_Grantham")
        for nm in ("Matrix_count", "Matrix_positions"):
            xw.sheets[nm].conditional_formatting.add("B2:U21", ColorScaleRule(start_type="num", start_value=0, start_color="FFFFFF",
                         mid_type="percentile", mid_value=90, mid_color="FFB74D", end_type="max", end_color="C62828"))
        xw.sheets["Matrix_Grantham"].conditional_formatting.add("B2:U21", ColorScaleRule(start_type="num", start_value=0, start_color="2E7D32",
                         mid_type="num", mid_value=100, mid_color="FFF59D", end_type="num", end_value=200, end_color="C62828"))
        for name, wsx in xw.sheets.items():
            if name != "Summary": style_sheet(wsx)
            colour_category_column(wsx, "Grantham_category"); colour_category_column(wsx, "Grantham category")


# ------------------------------ driver ------------------------------
def list_inputs(path):
    files = []
    if isinstance(path, (list, tuple)):
        files = [str(p) for p in path]
    elif os.path.isdir(path):
        for f in sorted(os.listdir(path)):
            if f.lower().endswith((".xlsx", ".xls", ".csv")) and not f.startswith(("~", ".")):
                files.append(os.path.join(path, f))
    else:
        files.append(path)
    items = []
    for f in files:
        base = os.path.splitext(os.path.basename(f))[0]
        if f.lower().endswith(".csv"):
            items.append((base, pd.read_csv(f)))
        else:
            sheets = pd.read_excel(f, sheet_name=None)
            for sh, d in sheets.items():
                items.append((base if len(sheets) == 1 else f"{base}_{sh}", d))
    return items


def run(input_path, outdir):
    os.makedirs(outdir, exist_ok=True)
    allrows = []
    for gene, df in list_inputs(input_path):
        gene = re.sub(r"[^A-Za-z0-9_\-]", "_", gene)
        print(f"== {gene}")
        t, raw = read_gene_table(df, gene)
        u, info = clean_and_collapse(t, gene)
        if len(u) == 0:
            print("  no usable substitutions, skipped"); continue
        u = add_scores(u)
        res = analyse(u)
        cnt, pos, gr = matrices(u)
        # original rows + scores (for your chart)
        raw_out = raw.copy()
        parsed = t.set_index("_row")
        raw_out["Ref AA"] = parsed.Ref.reindex(range(len(raw))).values; raw_out["Alt AA"] = parsed.Alt.reindex(range(len(raw))).values
        raw_out["Grantham score"] = [GRANTHAM[a][b] if a in GRANTHAM and b in GRANTHAM[a] else np.nan for a, b in zip(raw_out["Ref AA"], raw_out["Alt AA"])]
        raw_out["Grantham category"] = raw_out["Grantham score"].map(lambda d: gcat(d) if pd.notna(d) and d > 0 else None)
        gdir = os.path.join(outdir, gene); os.makedirs(gdir, exist_ok=True)
        write_gene_workbook(os.path.join(gdir, f"{gene}_Grantham_report.xlsx"), gene, info, res, u, raw_out, cnt, pos, gr)
        draw_figures(u, cnt, gene, gdir)
        txt = results_text(gene, info, res, u)
        open(os.path.join(gdir, f"{gene}_results_text.txt"), "w", encoding="utf-8").write(txt)
        print(f"  {info['unique_substitutions']} unique substitutions; agreement with Grantham: {res['agreement'].iloc[0].agreement_pct:.1f}%")
        m = res["metrics"].set_index("metric"); ct = res["crosstab"]; ag = res["agreement"].set_index("criterion")
        row = dict(Gene=gene, **info)
        for i, c in enumerate(CATS):
            row[f"n {c}"] = int(ct.loc[i, "n"]); row[f"% {c}"] = round(ct.loc[i, "% of all unique"], 1)
            row[f"% class-radical within {c}"] = round(ct.loc[i, "Class-radical (%)"], 1)
        for met in ("Grantham",):
            if met in m.index and "cons_mean" in m.columns:
                row[f"{met} mean cons"] = round(m.loc[met, "cons_mean"], 2); row[f"{met} mean rad"] = round(m.loc[met, "rad_mean"], 2)
                row[f"{met} MWU p"] = m.loc[met, "mannwhitney_p"]; row[f"{met} AUC"] = round(m.loc[met, "AUC"], 3)
        for crit, r in ag.iterrows():
            row[f"Agreement % {crit}"] = round(r.agreement_pct, 1); row[f"Weighted agreement % {crit}"] = round(r.agreement_weighted_pct, 1)
            row[f"Kappa {crit}"] = round(r.kappa, 2)
        allrows.append(row)
    if allrows:
        pd.DataFrame(allrows).to_excel(os.path.join(outdir, "ALL_GENES_summary.xlsx"), index=False)
        print(f"\nDone. See {outdir}/ALL_GENES_summary.xlsx and one sub-folder per gene.")


def in_colab():
    try:
        import google.colab  # noqa: F401
        return True
    except ImportError:
        return False


def upload_box():
    """Open an upload / file-selection box. Returns (list_of_files, default_output_folder)."""
    if in_colab():
        from google.colab import files
        print("Click 'Choose files' below and select your gene file(s) (.xlsx / .csv)...")
        up = files.upload()
        names = [n for n in up if n.lower().endswith((".xlsx", ".xls", ".csv"))]
        return names, os.path.join(os.getcwd(), "results")
    try:
        import tkinter as tk
        from tkinter import filedialog
        root = tk.Tk(); root.withdraw(); root.attributes("-topmost", True)
        paths = filedialog.askopenfilenames(title="Select your gene file(s) (.xlsx / .csv) - hold Ctrl/Shift for several",
                                            filetypes=[("Data files", "*.xlsx *.xls *.csv"), ("All files", "*.*")])
        root.destroy()
        paths = list(paths)
        base = os.path.dirname(os.path.abspath(paths[0])) if paths else os.getcwd()
        return paths, os.path.join(base, "results")
    except Exception as e:
        print(f"Could not open a file-selection window ({e.__class__.__name__}).\n"
              "Please type the path of your data folder or file instead:")
        p = input("Path: ").strip().strip('"').strip("'")
        base = p if os.path.isdir(p) else os.path.dirname(os.path.abspath(p))
        return p, os.path.join(base, "results")


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--input", default=None, help="optional: folder or file (overrides the box)")
    ap.add_argument("--outdir", default=None, help="optional: output folder (default: 'results' next to the input)")
    a, _ = ap.parse_known_args()      # parse_known_args so it also works inside Jupyter/Colab
    inp = a.input or INPUT_PATH
    if inp:
        base = inp if os.path.isdir(inp) else os.path.dirname(os.path.abspath(inp))
        outdir = a.outdir or os.path.join(base, "results")
    else:
        inp, outdir = upload_box()
        outdir = a.outdir or outdir
    if not inp or (not isinstance(inp, list) and not os.path.exists(inp)):
        sys.exit("No input selected / found. Run again and choose your file(s), or set INPUT_PATH at the top.")
    run(inp, outdir)
    if in_colab():
        import shutil
        from google.colab import files
        shutil.make_archive("results", "zip", outdir)
        files.download("results.zip")